import io
import re
from typing import Optional


class UnreadableDocument(Exception):
    pass


def extract_bytes(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return extract_pdf(data)
    if name.endswith(".docx"):
        return extract_docx(data)
    if name.endswith(".txt") or name.endswith(".md"):
        return data.decode("utf-8", errors="replace")
    raise UnreadableDocument("Unsupported file type. Use PDF, DOCX, or TXT.")


def extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise UnreadableDocument("pypdf is not installed") from e
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"[[page={i}]]\n{text}")
    raw = "\n\n".join(pages).strip()
    if len(re.sub(r"\s+", "", raw)) < 40:
        raise UnreadableDocument(
            "This PDF has no extractable text. It may be a scan. Paste the scope section or export to Word."
        )
    return raw


def extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as e:
        raise UnreadableDocument("python-docx is not installed") from e
    doc = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    if len(re.sub(r"\s+", "", text)) < 40:
        raise UnreadableDocument("This Word file has no extractable text.")
    return text


def page_hint(raw: str, excerpt: str) -> Optional[int]:
    idx = raw.find(excerpt[:80]) if excerpt else -1
    if idx < 0:
        return None
    prefix = raw[:idx]
    pages = re.findall(r"\[\[page=(\d+)\]\]", prefix)
    return int(pages[-1]) if pages else None
